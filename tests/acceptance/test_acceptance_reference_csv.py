"""Prueba de aceptación de extremo a extremo sobre el CSV de referencia del proyecto."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import responses
from dirty_equals import IsPartialDict
from openpyxl import load_workbook

from rpa_cuentas.application.use_cases.audit_summary import EXIT_PARTIAL_FAILURE, EXIT_SUCCESS
from rpa_cuentas.main import main
from tests.conftest import make_test_env

BASE_URL = "https://jsonplaceholder.typicode.com"
EXPECTED_UNIQUE_ORDER: list[int] = [7, 12, 23, 45, 99, 101, 3, 58, 88, 150, 64, 31, 42, 76, 19]
NOT_FOUND_IDS: frozenset[int] = frozenset({101, 150})
REFERENCE_CSV = Path(__file__).resolve().parents[1] / "fixtures" / "cuentas_reference.csv"


def _register_reference_http_mocks(failing_account_id: int | None = None) -> None:
    """Registra las respuestas simuladas de JSONPlaceholder para las 15 cuentas únicas."""
    for account_id in EXPECTED_UNIQUE_ORDER:
        url = f"{BASE_URL}/posts/{account_id}"
        if account_id == failing_account_id:
            for _ in range(4):
                responses.add(responses.GET, url, status=503)
        elif account_id in NOT_FOUND_IDS:
            responses.add(responses.GET, url, json={}, status=404)
        else:
            payload = {
                "userId": (account_id % 10) + 1, "id": account_id,
                "title": f"Título cuenta {account_id}",
                "body": f"Contenido validado de la cuenta {account_id}",
            }  # fmt: skip
            responses.add(responses.GET, url, json=payload, status=200)


def _read_excel_sheets(
    excel_path: Path,
) -> tuple[list[tuple[Any, ...]], list[tuple[Any, ...]], dict[str, Any]]:
    """Lee las hojas Resultados, Incidencias y Resumen del archivo Excel generado."""
    wb = load_workbook(excel_path)
    assert wb.sheetnames == ["Resultados", "Incidencias", "Resumen"]
    resultados_rows = list(wb["Resultados"].iter_rows(values_only=True))
    incidencias_rows = list(wb["Incidencias"].iter_rows(values_only=True))
    resumen_rows = list(wb["Resumen"].iter_rows(values_only=True))
    wb.close()
    return resultados_rows, incidencias_rows, {str(k): v for k, v in resumen_rows[1:]}


@responses.activate
def test_acceptance_reference_csv_first_and_second_idempotent_runs(tmp_path: Path) -> None:
    """Verifica todos los números de aceptación en la 1ª corrida y 0 consultas en la 2ª corrida."""
    env = make_test_env(tmp_path)
    excel_path, lock_path = Path(env["OUTPUT_EXCEL_PATH"]), Path(env["LOCK_FILE_PATH"])
    _register_reference_http_mocks()

    exit_code_1 = main(["run", "--input", str(REFERENCE_CSV)], environ=env, sleep_fn=lambda _: None)
    assert (exit_code_1, len(responses.calls)) == (EXIT_SUCCESS, 15)
    assert excel_path.exists() and not lock_path.exists()

    resultados_1, incidencias_1, resumen_1 = _read_excel_sheets(excel_path)

    assert resumen_1 == IsPartialDict(
        lineas_leidas=21, lineas_vacias=1, invalidas_no_entero=1, invalidas_no_positivo=1,
        total_incidencias=3, filas_validas=18, duplicados_colapsados=3, cuentas_unicas=15,
        consultadas_en_ejecucion=15, reutilizadas_de_estado=0, resueltas=13, no_existe=2,
        fallos_tecnicos=0, codigo_salida=0,
    )  # fmt: skip

    assert incidencias_1[1:] == [
        (9, "abc", "NO_ENTERO"), (12, None, "LINEA_VACIA"), (21, "0", "NO_POSITIVO"),
    ]  # fmt: skip

    data_rows_1 = resultados_1[1:]
    assert [row[0] for row in data_rows_1] == EXPECTED_UNIQUE_ORDER

    by_id_1 = {int(row[0]): row for row in data_rows_1}
    assert (by_id_1[7][1], by_id_1[7][3]) == ("RESUELTA", 3)
    assert (by_id_1[12][1], by_id_1[12][3]) == ("RESUELTA", 2)
    assert by_id_1[101][1:3] == ("NO_EXISTE", "RECURSO_NO_ENCONTRADO")
    assert by_id_1[150][1:3] == ("NO_EXISTE", "RECURSO_NO_ENCONTRADO")

    # Segunda ejecución inmediata: 0 consultas HTTP nuevas y mismo contenido en Resultados
    responses.reset()
    exit_code_2 = main(["run", "--input", str(REFERENCE_CSV)], environ=env, sleep_fn=lambda _: None)
    assert exit_code_2 == EXIT_SUCCESS and len(responses.calls) == 0

    resultados_2, incidencias_2, resumen_2 = _read_excel_sheets(excel_path)
    assert resumen_2 == IsPartialDict(
        consultadas_en_ejecucion=0, reutilizadas_de_estado=15, resueltas=13,
        no_existe=2, fallos_tecnicos=0, codigo_salida=0,
    )  # fmt: skip
    assert (resultados_2, incidencias_2) == (resultados_1, incidencias_1)


@responses.activate
def test_acceptance_reference_csv_technical_failure_and_selective_recovery(tmp_path: Path) -> None:
    """Verifica código 1 ante fallo técnico en 45 y re-consulta exclusiva de 45 al recuperar."""
    env = make_test_env(tmp_path)
    excel_path = Path(env["OUTPUT_EXCEL_PATH"])
    _register_reference_http_mocks(failing_account_id=45)

    exit_code_fail = main(
        argv=["run", "--input", str(REFERENCE_CSV)], environ=env, sleep_fn=lambda _: None
    )
    assert (exit_code_fail, len(responses.calls)) == (EXIT_PARTIAL_FAILURE, 18)

    resultados_fail, _, resumen_fail = _read_excel_sheets(excel_path)
    assert resumen_fail == IsPartialDict(
        resueltas=12, no_existe=2, fallos_tecnicos=1, codigo_salida=1
    )

    by_id_fail = {int(row[0]): row for row in resultados_fail[1:]}
    assert (by_id_fail[45][1], by_id_fail[45][2], by_id_fail[45][7]) == (
        "FALLO_TECNICO", "HTTP_5XX_AGOTADO: 503", 4,
    )  # fmt: skip

    # Relanzamiento con servicio recuperado: solo consulta la cuenta 45
    responses.reset()
    responses.add(
        responses.GET,
        f"{BASE_URL}/posts/45",
        json={"userId": 6, "id": 45, "title": "Título cuenta 45", "body": "Contenido 45"},
        status=200,
    )

    exit_code_recovered = main(
        argv=["run", "--input", str(REFERENCE_CSV)], environ=env, sleep_fn=lambda _: None
    )
    assert exit_code_recovered == EXIT_SUCCESS and len(responses.calls) == 1
    assert responses.calls[0].request.url == f"{BASE_URL}/posts/45"

    resultados_ok, _, resumen_ok = _read_excel_sheets(excel_path)
    assert resumen_ok == IsPartialDict(
        consultadas_en_ejecucion=1, reutilizadas_de_estado=14, resueltas=13,
        no_existe=2, fallos_tecnicos=0, codigo_salida=0,
    )  # fmt: skip

    by_id_ok = {int(row[0]): row for row in resultados_ok[1:]}
    assert (by_id_ok[45][1], by_id_ok[45][7]) == ("RESUELTA", 1)
