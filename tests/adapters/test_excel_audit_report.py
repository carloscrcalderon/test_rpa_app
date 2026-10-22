"""Pruebas unitarias para el adaptador de reporte ExcelReportWriter."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from rpa_cuentas.adapters.outbound.file_io.excel_audit_report import ExcelReportWriter
from rpa_cuentas.application.ports.outbound_ports import ReportPublicationError
from rpa_cuentas.application.use_cases.audit_summary import BatchSummary
from rpa_cuentas.domain import AccountStatus, IncidentReason, InputIncident
from tests.conftest import make_row, make_summary


def _sample_summary() -> BatchSummary:
    """Crea un BatchSummary de prueba."""
    return make_summary(
        run_id="run-test-01", lines_read=21, empty_lines=1, total_incidents=3,
        valid_lines=18, duplicates_collapsed=3, unique_accounts=15, queried_in_run=15,
        resolved_count=12, not_found_count=2, technical_failure_count=1, exit_code=1,
        invalid_lines_by_reason={"LINEA_VACIA": 1, "NO_ENTERO": 1, "NO_POSITIVO": 1},
        failures_by_cause={"TIMEOUT_AGOTADO": 1},
    )  # fmt: skip


def test_excel_report_writer_generates_three_sheets_with_expected_schema(tmp_path: Path) -> None:
    """Verifica las hojas Resultados, Incidencias y Resumen y ausencia de temporales residuales."""
    output_path = tmp_path / "output" / "consolidado_cuentas.xlsx"
    writer = ExcelReportWriter(output_path)

    rows = [
        make_row(7, apariciones=3, user_id=1, titulo="Cuenta 7", cuerpo="Detalle 7"),
        make_row(
            101,
            AccountStatus.NO_EXISTE,
            "RECURSO_NO_ENCONTRADO",
            fecha_consulta="2026-10-21T03:00:01Z",
        ),
        make_row(
            45,
            AccountStatus.FALLO_TECNICO,
            "TIMEOUT_AGOTADO",
            intentos=4,
            fecha_consulta="2026-10-21T03:00:05Z",
        ),
    ]
    incidents = [
        InputIncident(9, "abc", IncidentReason.NO_ENTERO),
        InputIncident(12, "", IncidentReason.LINEA_VACIA),
    ]
    writer.publish_report(rows=rows, incidents=incidents, summary=_sample_summary())

    assert output_path.exists() and list(output_path.parent.iterdir()) == [output_path]

    wb = load_workbook(output_path)
    assert wb.sheetnames == ["Resultados", "Incidencias", "Resumen"]

    res_rows = list(wb["Resultados"].iter_rows(values_only=True))
    assert res_rows[0] == (
        "id_cuenta", "estado", "causa", "apariciones", "user_id",
        "titulo", "cuerpo", "intentos", "fecha_consulta",
    )  # fmt: skip
    assert res_rows[1] == (
        7, "RESUELTA", None, 3, 1, "Cuenta 7", "Detalle 7", 1, "2026-10-21T03:00:00Z"
    )  # fmt: skip
    assert res_rows[2] == (
        101, "NO_EXISTE", "RECURSO_NO_ENCONTRADO", 1, None, None, None, 1, "2026-10-21T03:00:01Z"
    )  # fmt: skip
    assert res_rows[3] == (
        45, "FALLO_TECNICO", "TIMEOUT_AGOTADO", 1, None, None, None, 4, "2026-10-21T03:00:05Z"
    )  # fmt: skip

    inc_rows = list(wb["Incidencias"].iter_rows(values_only=True))
    assert inc_rows == [
        ("linea", "valor_original", "motivo"),
        (9, "abc", "NO_ENTERO"),
        (12, None, "LINEA_VACIA"),
    ]

    sum_rows = dict(list(wb["Resumen"].iter_rows(values_only=True))[1:])
    assert sum_rows == {
        "run_id": "run-test-01",
        "lineas_leidas": 21,
        "lineas_vacias": 1,
        "invalidas_no_entero": 1,
        "invalidas_no_positivo": 1,
        "total_incidencias": 3,
        "filas_validas": 18,
        "duplicados_colapsados": 3,
        "cuentas_unicas": 15,
        "consultadas_en_ejecucion": 15,
        "reutilizadas_de_estado": 0,
        "resueltas": 12,
        "no_existe": 2,
        "fallos_tecnicos": 1,
        "detalle_fallos_por_causa": "TIMEOUT_AGOTADO: 1",
        "codigo_salida": 1,
    }
    wb.close()


def test_excel_report_writer_preserves_previous_report_and_cleans_temp_on_replace_failure(
    tmp_path: Path,
) -> None:
    """Verifica que si falla os.replace, el destino previo quede intacto y sin temporales."""
    output_path = tmp_path / "output" / "consolidado_cuentas.xlsx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    previous_bytes = b"reporte_previo_intacto"
    output_path.write_bytes(previous_bytes)

    def failing_replace(_src: Path | str, _dst: Path | str) -> None:
        raise OSError("Simulated disk error during atomic replace")

    writer = ExcelReportWriter(output_path, replace_fn=failing_replace)

    with pytest.raises(ReportPublicationError, match="Simulated disk error"):
        writer.publish_report(rows=[], incidents=[], summary=_sample_summary())

    assert output_path.read_bytes() == previous_bytes
    assert list(output_path.parent.iterdir()) == [output_path]
